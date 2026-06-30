import os
import json
import time
import re
import requests
import pandas as pd
from datetime import datetime, date
from bs4 import BeautifulSoup
from playwright.sync_api import sync_playwright
from openai import OpenAI
from openpyxl import load_workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from urllib.parse import urlparse
from dotenv import load_dotenv

load_dotenv()

# ==================================================
# CONFIG
# ==================================================
START_URL = "https://www.studyabroad.pk/scholarships/"
LOAD_MORE_CLICKS = 10
BATCH_SIZE = 5
OUTPUT_EXCEL = "Data.xlsx"
SHEET_NAME = "study_abroad"
TODAY = date.today()

# ===========================
# OPENAI API KEY
# ===========================
OPENAI_API_KEY = os.environ["OPENAI_API_KEY"]

# ==================================================
# OPENAI CLIENT
# ==================================================
client = OpenAI(api_key=OPENAI_API_KEY)


# ==================================================
# HELPER — Safely convert any value to a plain string
# ==================================================
def safe_str(val) -> str:
    """Convert lists/tuples to comma-separated string; everything else to str."""
    if val is None:
        return ""
    if isinstance(val, (list, tuple)):
        return ", ".join(str(v).strip() for v in val if v)
    return str(val).strip()


# ==================================================
# DEADLINE FILTER & PARSING
# ==================================================
def parse_deadline(deadline_str: str) -> date | None:
    """Try to parse a deadline string into a date object."""
    if not deadline_str or not deadline_str.strip():
        return None

    deadline_str = deadline_str.strip()

    formats = [
        "%d %B %Y",
        "%d %b %Y",
        "%B %d, %Y",
        "%b %d, %Y",
        "%Y-%m-%d",
        "%d/%m/%Y",
        "%m/%d/%Y",
        "%d-%b-%y",
        "%d-%b-%Y",
        "%d-%B-%Y",
        "%d %B, %Y",
        "%d %b, %Y",
        "%b %d %Y",
        "%B %d %Y",
        "%d/%m/%y",
        "%m/%d/%y",
    ]

    for fmt in formats:
        try:
            return datetime.strptime(deadline_str, fmt).date()
        except ValueError:
            continue

    m = re.search(
        r"(\d{1,2})\s+(January|February|March|April|May|June|July|August|"
        r"September|October|November|December)\s+(\d{4})",
        deadline_str, re.IGNORECASE
    )
    if m:
        try:
            return datetime.strptime(f"{m.group(1)} {m.group(2)} {m.group(3)}", "%d %B %Y").date()
        except ValueError:
            pass

    m = re.search(
        r"(January|February|March|April|May|June|July|August|"
        r"September|October|November|December)\s+(\d{1,2}),?\s+(\d{4})",
        deadline_str, re.IGNORECASE
    )
    if m:
        try:
            return datetime.strptime(f"{m.group(2)} {m.group(1)} {m.group(3)}", "%d %B %Y").date()
        except ValueError:
            pass

    return None


def is_deadline_passed(deadline_str: str) -> bool:
    """Return True if the deadline is before today. If unparseable, keep the scholarship."""
    d = parse_deadline(deadline_str)
    if d is None:
        return False
    return d < TODAY


# ==================================================
# EXCEL APPEND HELPER
# ==================================================
def append_df_to_excel(filename: str, df: pd.DataFrame, sheet_name: str):
    """
    Append a DataFrame to an existing excel file / sheet.
    If file doesn't exist -> create it and write header + data.
    If sheet doesn't exist -> create sheet with header + data.
    If sheet exists -> append rows without headers.
    """
    # ---- FIX: Ensure every column is a plain scalar (no lists) ----
    for col in df.columns:
        df[col] = df[col].apply(safe_str)

    if not os.path.exists(filename):
        with pd.ExcelWriter(filename, engine="openpyxl") as writer:
            df.to_excel(writer, sheet_name=sheet_name, index=False)
        return

    wb = load_workbook(filename)
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for r in dataframe_to_rows(df, index=False, header=False):
            ws.append(r)
        wb.save(filename)
    else:
        with pd.ExcelWriter(filename, engine="openpyxl", mode="a") as writer:
            df.to_excel(writer, sheet_name=sheet_name, index=False)


# ==================================================
# STEP 1 — DISCOVER SCHOLARSHIP LINKS (WITH LOAD MORE)
# ==================================================
print("🌐 Launching browser and loading scholarships...")

scholarships = []
allowed_hosts = {"studyabroad.pk", "www.studyabroad.pk"}

with sync_playwright() as p:
    browser = p.chromium.launch(headless=True)
    page = browser.new_page()
    page.goto(START_URL, wait_until="domcontentloaded", timeout=60000)

    for i in range(LOAD_MORE_CLICKS):
        try:
            page.locator("#loadMoreScholarshipsBtn").scroll_into_view_if_needed()
            page.click("#loadMoreScholarshipsBtn")
            time.sleep(1.5)
            print(f"✅ Load More clicked {i+1}/{LOAD_MORE_CLICKS}")
        except Exception:
            print("⚠ Load More button not found or no more items.")
            break

    html = page.content()
    soup = BeautifulSoup(html, "html.parser")

    for row in soup.select("#list-scholarship table tbody tr"):
        link_el = row.select_one(".uni-name a")
        if not link_el:
            continue

        title = link_el.get_text(strip=True)
        href = link_el.get("href", "") or ""
        href = href.strip()

        if not href or href.lower().startswith("javascript:") or href == "#":
            continue

        link = href

        m = re.search(r"https?://[^\s\"']+", link)
        if m:
            candidate = m.group(0)
            parsed_candidate = urlparse(candidate)
            host = (parsed_candidate.netloc or "").lower().split(":")[0]
            if host in allowed_hosts:
                link = candidate
            else:
                print(f"⚠ Link not valid (external embedded): {candidate} — skipping")
                continue
        elif link.startswith("/"):
            link = "https://www.studyabroad.pk" + link
        else:
            parsed = urlparse(link)
            host = (parsed.netloc or "").lower().split(":")[0]
            if host:
                if host not in allowed_hosts:
                    print(f"⚠ Link not valid (external): {link} — skipping")
                    continue
            else:
                link = "https://www.studyabroad.pk/" + link.lstrip("/")

        parsed_final = urlparse(link)
        final_host = (parsed_final.netloc or "").lower().split(":")[0]
        if final_host not in allowed_hosts:
            print(f"⚠ Final link not valid: {link} — skipping")
            continue

        if title and link:
            scholarships.append({"name": title, "link": link})
            print(f"✔ {title} -> {link}")

    browser.close()

print(f"\n📦 Found {len(scholarships)} scholarships\n")

# ==================================================
# STEP 1.5 — LOAD EXISTING LINKS FROM EXCEL (skip duplicates)
# ==================================================
existing_links = set()
if os.path.exists(OUTPUT_EXCEL):
    try:
        df_existing = pd.read_excel(OUTPUT_EXCEL, sheet_name=SHEET_NAME)
        if "Scholarship Link" in df_existing.columns:
            existing_links = set(df_existing["Scholarship Link"].astype(str).str.strip().dropna().tolist())
        else:
            if df_existing.shape[1] >= 2:
                colname = df_existing.columns[1]
                existing_links = set(df_existing[colname].astype(str).str.strip().dropna().tolist())
    except Exception as e:
        print("⚠ Could not read existing excel sheet:", e)

print(f"🔁 Loaded {len(existing_links)} existing links from {OUTPUT_EXCEL} (sheet: {SHEET_NAME})")

# ==================================================
# STEP 2 — SCRAPE + AI STRUCTURE
# ==================================================
print("🧭 Launching browser...")

results_batch = []
success_count = 0
skipped_count = 0

with sync_playwright() as p:
    browser = p.chromium.launch(headless=True)
    page = browser.new_page()

    for s in scholarships:
        if s["link"] in existing_links:
            print(f"  ⏭ Already in sheet — skipping: {s['link']}")
            skipped_count += 1
            continue

        print(f"\n🔍 Scraping: {s['name']}")

        try:
            page.goto(s["link"], wait_until="domcontentloaded", timeout=60000)
        except Exception as e:
            print("⚠ Page load failed:", e)
            continue

        try:
            page.evaluate("window.scrollTo(0, document.body.scrollHeight)")
            time.sleep(1)
        except:
            pass

        text = page.evaluate("""
        () => {
            const main =
                document.querySelector('.sec') ||
                document.querySelector('.content') ||
                document.body;
            return main ? main.innerText : document.body.innerText;
        }
        """)

        try:
            official_website = page.evaluate("""
            () => {
                const anchors = Array.from(document.querySelectorAll('a'));
                const patterns = [/^\\s*apply\\s*now\\s*$/i, /^\\s*applynow\\s*$/i, /^\\s*apply\\s*$/i];
                for (const a of anchors) {
                    try {
                        const text = (a.innerText || '').trim();
                        if (!text) continue;
                        const rect = a.getBoundingClientRect();
                        const style = window.getComputedStyle(a);
                        const visible = rect.width > 0 && rect.height > 0 && style.visibility !== 'hidden' && style.display !== 'none' && style.opacity !== '0';
                        if (!visible) continue;
                        for (const p of patterns) {
                            if (p.test(text)) {
                                const href = a.href || '';
                                if (!href) continue;
                                if (href.startsWith('javascript:') || href === location.href || href === '#') continue;
                                return href;
                            }
                        }
                    } catch (e) {}
                }
                return '';
            }
            """)
        except Exception:
            official_website = ""

        # OPENAI STRUCTURING
        # Prompt explicitly asks for string values (not arrays) to prevent the list bug
        prompt = f"""
You are a data extraction engine.

STRICT RULES:
- Return ONLY valid JSON
- No explanations, no markdown, no extra text
- All values MUST be plain strings (never arrays or lists)
- If multiple values exist for a field, join them with a comma into one string

JSON schema:
{{
  "scholarship_type": "",
  "deadline": "",
  "host_country": "",
  "degree_type": "",
  "field_of_study": ""
}}

Text:
{text}
"""

        try:
            response = client.responses.create(
                model="gpt-4.1",
                input=prompt,
                temperature=0
            )

            chunks = []
            for item in response.output:
                for block in item.content:
                    if hasattr(block, "text"):
                        chunks.append(block.text)

            raw = "".join(chunks).strip()
            # Strip markdown fences if present
            raw = re.sub(r"^```[a-z]*\n?", "", raw, flags=re.IGNORECASE)
            raw = re.sub(r"\n?```$", "", raw)
            raw = raw.strip()

            data = json.loads(raw)

            if isinstance(data, list):
                data = data[0] if data else {}
            elif not isinstance(data, dict):
                data = {}

        except Exception as e:
            print("⚠ OpenAI failed:", e)
            data = {
                "scholarship_type": "",
                "deadline": "",
                "host_country": "",
                "degree_type": "",
                "field_of_study": ""
            }

        # --- CHECK DEADLINE BEFORE ADDING ---
        deadline_raw = safe_str(data.get("deadline", ""))
        if is_deadline_passed(deadline_raw):
            print(f"  ⏭ SKIPPED (deadline passed: {deadline_raw})")
            skipped_count += 1
            continue

        parsed_date = parse_deadline(deadline_raw)
        deadline_standard = parsed_date.isoformat() if parsed_date else ""

        record = {
            "Scholarship Name": safe_str(s["name"]),
            "Scholarship Link": safe_str(s["link"]),
            "Official Website":  safe_str(official_website),
            "Scholarship Type":  safe_str(data.get("scholarship_type", "")),
            "Deadline":          deadline_standard,
            "Host Country":      safe_str(data.get("host_country", "")),
            "Degree Type":       safe_str(data.get("degree_type", "")),
            "Field of Study":    safe_str(data.get("field_of_study", "")),
        }

        results_batch.append(record)
        success_count += 1
        print("  ✅ Added")

        if success_count % BATCH_SIZE == 0:
            df_batch = pd.DataFrame(results_batch)
            try:
                append_df_to_excel(OUTPUT_EXCEL, df_batch, SHEET_NAME)
                print(f"✅ Wrote batch of {len(results_batch)} to {OUTPUT_EXCEL} (sheet: {SHEET_NAME})")
                for ln in df_batch["Scholarship Link"].astype(str).str.strip().dropna().tolist():
                    existing_links.add(ln)
                results_batch.clear()
            except Exception as e:
                print("⚠ Failed to append batch to Excel:", e)

    browser.close()

# Write any remaining results
if results_batch:
    df_batch = pd.DataFrame(results_batch)
    try:
        append_df_to_excel(OUTPUT_EXCEL, df_batch, SHEET_NAME)
        print(f"✅ Wrote final batch of {len(results_batch)} to {OUTPUT_EXCEL} (sheet: {SHEET_NAME})")
    except Exception as e:
        print(f"⚠ Failed to append final batch to Excel: {e}")
else:
    print("\n⚠ No remaining results to write.")

print(f"\n📊 Saved: {success_count} | Skipped (expired or existing): {skipped_count}")